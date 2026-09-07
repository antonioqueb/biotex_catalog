"""Read incremental catalog sheets without confusing legacy fields with stock."""
import hashlib
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return unicodedata.normalize('NFC', str(value)).strip()


def useful(value):
    value = text(value)
    return value if value.upper() not in ('', 'N/A', 'N/D', 'NA', 'S/R') else False


def value(row, index):
    values = row['values']
    return values[index] if index < len(values) else None


def read_products(filename):
    path = Path(filename)
    evaluated = load_workbook(path, read_only=True, data_only=True)
    original = load_workbook(path, read_only=True, data_only=False)
    try:
        if set(evaluated.sheetnames) != {'RECLASIFICACION', 'NUEVAS'}:
            raise ValueError('Se requieren exactamente las hojas RECLASIFICACION y NUEVAS.')
        rows, columns, counts = [], {}, {}
        for name in evaluated.sheetnames:
            cells = list(evaluated[name].values)
            formulas = list(original[name].values)
            width = 56 if name == 'RECLASIFICACION' else 21
            headers = list(cells[0])
            if len(headers) != width or headers[5:7] != ['CLAVE', 'DESCRIPCION']:
                raise ValueError('Encabezados no reconocidos: ' + name)
            if width == 56 and headers[21:24] != ['CLAVE', 'CLAVE ALTERNA', 'DESCRIPCION']:
                raise ValueError('No se reconoce el bloque anterior de RECLASIFICACION.')
            columns[name] = headers
            counts[name] = 0
            for index, values in enumerate(cells[1:], 1):
                if not any(text(v) for v in values):
                    continue
                row = {'worksheet': name, 'worksheet_row': index + 1, 'values': list(values),
                       'formulas': [v if isinstance(v, str) and v.startswith('=') else None
                                    for v in formulas[index]]}
                if not useful(value(row, 6)) and not useful(value(row, 23)):
                    raise ValueError('Producto sin descripción: %s, fila %s' % (name, index + 1))
                for column in (5, 6, 14, 21, 23):
                    if text(value(row, column)).startswith(('#REF!', '#VALUE!', '#DIV/0!', '#N/A')):
                        raise ValueError('Identidad con error de Excel: %s, fila %s' % (name, index + 1))
                rows.append(row)
                counts[name] += 1
        return {'filename': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                'columns': columns, 'rows': rows, 'sheet_counts': counts}
    finally:
        evaluated.close()
        original.close()


def identifiers(row):
    result = []
    for index, namespace in ((5, 'reference'), (21, 'reference'), (14, 'barcode')):
        item = useful(value(row, index))
        if item:
            result.append((namespace, item))
    return result


def plan_rows(data, products):
    """Names/manufacturer references alone never merge product identities."""
    index = defaultdict(set)
    for product in products:
        for field in ('code', 'legacy'):
            if useful(product.get(field)):
                index[('reference', useful(product[field]))].add(product['id'])
        for variant in product.get('variants') or []:
            for field, namespace in (('code', 'reference'), ('barcode', 'barcode')):
                if useful(variant.get(field)):
                    index[(namespace, useful(variant[field]))].add(product['id'])
        for row in (product.get('source') or {}).get('rows', []):
            for identifier in identifiers(row):
                index[identifier].add(product['id'])
    result = []
    for row in data['rows']:
        keys = identifiers(row)
        if not keys:
            raise ValueError('Falta identificador estable: %s, fila %s' % (row['worksheet'], row['worksheet_row']))
        matches = set().union(*(index[key] for key in keys))
        if len(matches) > 1:
            raise ValueError('Identificadores apuntan a varios productos: %s, fila %s' %
                             (row['worksheet'], row['worksheet_row']))
        if matches:
            match = matches.pop()
            if isinstance(match, str):
                raise ValueError('Identificador repetido entre altas nuevas: %s, fila %s' %
                                 (row['worksheet'], row['worksheet_row']))
            result.append({'row': row, 'status': 'existing', 'product_id': match})
        else:
            planned = '%s:%s' % (row['worksheet'], row['worksheet_row'])
            for key in keys:
                index[key].add(planned)
            result.append({'row': row, 'status': 'create'})
    return result
