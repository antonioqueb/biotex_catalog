"""Read and validate the customer's workbook without guessing formulas or losing columns."""
import hashlib
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

SHEETS = {
    'divisions': '01_DIM_DIVISION', 'groups': '02_DIM_GRUPO',
    'families': '03_DIM_FAMILIA', 'classifiers': '04_DIM_CLASIFICADOR',
    'family_classifiers': '05_REL_FAMILIA_CLAS', 'brands': '06_DIM_MARCA',
    'specialties': '07_DIM_ESPECIALIDAD', 'equipment': '08_DIM_EQUIPO',
    'subclasses': '09_DIM_SUBCLASE_MT', 'product_specialties': '21_REL_PROD_ESPECIALIDAD',
    'product_equipment': '22_REL_PROD_EQUIPO',
}


def read_workbook(path):
    path = Path(path)
    book = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    data = {key: [list(row) for row in list(book[sheet].values)[4:] if row[0] is not None]
            for key, sheet in SHEETS.items()}
    data['headers'] = list(list(book['20_REMAPEO'].values)[4])
    rows = list(book['20_REMAPEO'].values)
    raw = list(formulas['20_REMAPEO'].values)
    data['rows'] = [{'worksheet_row': i + 1, 'values': list(row), 'formulas': list(raw[i])}
                    for i, row in enumerate(rows) if i >= 5 and row[56] is not None]
    data['incidents'] = [list(row) for row in list(book['31_INCIDENCIAS'].values)[4:] if row[0] is not None]
    data['sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    data['filename'] = path.name
    divisions = {r[0] for r in data['divisions']}
    groups = {r[0]: r for r in data['groups']}
    families = {(r[1], r[2]): r for r in data['families']}
    classifiers = {(r[1], r[2]): r for r in data['classifiers']}
    brands = {r[0] for r in data['brands']}
    pairs = {(r[1], r[2], r[3]) for r in data['family_classifiers'] if r[5] == 'SI'}
    for key, size in [('divisions', len(divisions)), ('groups', len(groups)),
                      ('families', len(families)), ('classifiers', len(classifiers)), ('brands', len(brands))]:
        if len(data[key]) != size:
            raise ValueError('Duplicate dimension keys: ' + key)
    for row in data['groups']:
        if row[2] not in divisions:
            raise ValueError('Unknown division: ' + str(row))
    for group, family, classifier in pairs:
        if (group, family) not in families or (group, classifier) not in classifiers:
            raise ValueError('Invalid family/classifier relationship')
    by_key = defaultdict(list)
    origins = set()
    for row in data['rows']:
        v = row['values']
        if v[56] in origins:
            raise ValueError('Repeated source row')
        origins.add(v[56])
        group, family, classifier, brand = v[58:62]
        if (group not in groups or groups[group][2] != v[57]
                or (group, family, classifier) not in pairs or brand not in brands):
            raise ValueError('Invalid classification on source row %s' % v[56])
        expected = '%s-%s-%s-%s-%02d' % (group, brand, family, classifier, int(v[62]))
        if v[63] != expected:
            raise ValueError('Inconsistent proposed code on row %s' % v[56])
        by_key[v[63]].append(row)
    for key in ('product_specialties', 'product_equipment'):
        dimension = {r[0] for r in data['specialties' if key.endswith('specialties') else 'equipment']}
        for row in data[key]:
            if row[0] not in by_key or row[1] not in dimension:
                raise ValueError('Unknown relationship: ' + str(row))
    data['products'] = dict(by_key)
    data['summary'] = {key: len(rows) for key, rows in data.items() if key in SHEETS}
    data['summary'].update(source_rows=len(data['rows']), unique_products=len(by_key),
                           duplicate_keys=sum(len(rows) > 1 for rows in by_key.values()))
    book.close()
    formulas.close()
    return data
